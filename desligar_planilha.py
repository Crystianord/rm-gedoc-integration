"""
Desliga (terminate) no GEDOC os funcionários listados numa planilha .xlsx, usando
o MESMO fluxo/endpoint da integração normal (gedoc_client.terminate_internal_employee
-> PUT /employees/{id}/terminate-options, autenticado com bearer token).

Fluxo:
    1. Lê nome+CPF de um arquivo texto (nome<TAB>cpf por linha) — evita depender de
       match por nome no RM, que falhava para nomes com grafia diferente.
    2. Lê os nomes da planilha .xlsx (aba "Ausentes na Lista 2") e casa cada linha
       com o CPF correspondente do arquivo de nomes+CPFs.
    3. Carrega o cache do GEDOC (gedoc_cache.py) para achar o registro (_id) de
       cada CPF.
    4. Para cada um encontrado e ativo: chama terminate_internal_employee.
    5. Gera um log .csv com o resultado de cada linha (desligado / não encontrado /
       erro / já desligado).

Por segurança, roda em modo DRY-RUN por padrão — não chama a API de desligamento,
só mostra o que faria. Para executar de verdade, use --confirmar.

Uso:
    python desligar_planilha.py <planilha.xlsx> <nomes_cpfs.txt>                 # dry-run
    python desligar_planilha.py <planilha.xlsx> <nomes_cpfs.txt> --confirmar     # executa de verdade

Formato de nomes_cpfs.txt (uma linha por pessoa, nome e CPF separados por TAB):
    FULANO DE TAL	12345678900
"""
import re
import sys
import csv
import logging
import unicodedata
from datetime import datetime

import openpyxl

from gedoc_client import GEDOCClient
from gedoc_cache import cache

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)


def normalize(name: str) -> str:
    if not name:
        return ""
    n = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    return " ".join(n.strip().upper().split())


def normalize_cpf(cpf: str) -> str:
    return re.sub(r"\D", "", cpf or "")


def ler_nomes_planilha(caminho: str) -> list[str]:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb["Ausentes na Lista 2"] if "Ausentes na Lista 2" in wb.sheetnames else wb.active
    nomes = []
    for row in range(2, ws.max_row + 1):
        v = ws.cell(row=row, column=1).value
        if v:
            nomes.append(str(v).strip())
    return nomes


def ler_nomes_cpfs(caminho: str) -> dict[str, list[str]]:
    """Retorna nome_normalizado -> lista de CPFs (pode ter mais de um se houver homônimo)."""
    mapa: dict[str, list[str]] = {}
    with open(caminho, encoding="utf-8") as f:
        for linha in f:
            linha = linha.strip()
            if not linha or "\t" not in linha:
                continue
            nome, cpf = linha.split("\t", 1)
            nome_norm = normalize(nome)
            cpf_norm = normalize_cpf(cpf)
            mapa.setdefault(nome_norm, []).append(cpf_norm)
    return mapa


def main():
    args = sys.argv[1:]
    if len(args) < 2:
        print("Uso: python desligar_planilha.py <planilha.xlsx> <nomes_cpfs.txt> [--confirmar]")
        sys.exit(1)

    caminho_planilha = args[0]
    caminho_cpfs = args[1]
    dry_run = "--confirmar" not in args

    if dry_run:
        print("=== MODO DRY-RUN (nenhuma chamada real será feita ao GEDOC) ===")
        print("Para executar de verdade, rode novamente com --confirmar\n")
    else:
        print("=== MODO REAL — VAI DESLIGAR FUNCIONÁRIOS NO GEDOC DE VERDADE ===\n")

    nomes = ler_nomes_planilha(caminho_planilha)
    print(f"{len(nomes)} linhas lidas da planilha.")

    cpfs_por_nome = ler_nomes_cpfs(caminho_cpfs)
    print(f"{sum(len(v) for v in cpfs_por_nome.values())} pares nome/CPF carregados do arquivo de referência.")

    print("Carregando cache do GEDOC (employees, employees/internal)...")
    cache.preload()

    gedoc = GEDOCClient()

    resultados = []
    # controla quantas vezes um nome duplicado já foi consumido, para casar 1:1
    # com a ordem de aparição no arquivo de nomes/CPFs quando há homônimos.
    consumidos: dict[str, int] = {}

    for nome in nomes:
        nome_norm = normalize(nome)
        cpfs = cpfs_por_nome.get(nome_norm)

        if not cpfs:
            resultados.append({
                "nome": nome, "cpf": "", "hiring_type": "", "gedoc_id": "",
                "acao": "IGNORADO", "detalhe": "Nome não encontrado no arquivo de nomes/CPFs",
            })
            continue

        idx = consumidos.get(nome_norm, 0)
        cpf = cpfs[idx] if idx < len(cpfs) else cpfs[-1]
        consumidos[nome_norm] = idx + 1

        if not cpf:
            resultados.append({
                "nome": nome, "cpf": "", "hiring_type": "", "gedoc_id": "",
                "acao": "IGNORADO", "detalhe": "CPF vazio no arquivo de referência",
            })
            continue

        gedoc_emp = cache.find_employee_by_cpf(cpf)
        if not gedoc_emp:
            resultados.append({
                "nome": nome, "cpf": cpf, "hiring_type": "", "gedoc_id": "",
                "acao": "IGNORADO", "detalhe": "CPF não encontrado no GEDOC",
            })
            continue

        hiring_type = gedoc_emp.get("_hiring_type", "?")
        gedoc_id = gedoc_emp.get("_id", "")

        if gedoc_emp.get("terminated"):
            resultados.append({
                "nome": nome, "cpf": cpf, "hiring_type": hiring_type, "gedoc_id": gedoc_id,
                "acao": "IGNORADO", "detalhe": "Já estava desligado no GEDOC",
            })
            continue

        if dry_run:
            resultados.append({
                "nome": nome, "cpf": cpf, "hiring_type": hiring_type, "gedoc_id": gedoc_id,
                "acao": "SERIA DESLIGADO", "detalhe": "Dry-run — nenhuma chamada real feita",
            })
            continue

        # --- Execução real ---
        ok = gedoc.terminate_internal_employee(gedoc_id, cpf)
        if ok:
            resultados.append({
                "nome": nome, "cpf": cpf, "hiring_type": hiring_type, "gedoc_id": gedoc_id,
                "acao": "DESLIGADO", "detalhe": "",
            })
        else:
            resultados.append({
                "nome": nome, "cpf": cpf, "hiring_type": hiring_type, "gedoc_id": gedoc_id,
                "acao": "ERRO", "detalhe": gedoc.last_error,
            })

    # --- Log CSV ---
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = f"log_desligamento_{timestamp}.csv"
    with open(log_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["nome", "cpf", "hiring_type", "gedoc_id", "acao", "detalhe"])
        writer.writeheader()
        writer.writerows(resultados)

    contagem: dict[str, int] = {}
    for r in resultados:
        contagem[r["acao"]] = contagem.get(r["acao"], 0) + 1

    print("\n=== Resumo ===")
    for acao, qtd in sorted(contagem.items()):
        print(f"  {acao}: {qtd}")
    print(f"\nLog salvo em: {log_path}")


if __name__ == "__main__":
    main()
