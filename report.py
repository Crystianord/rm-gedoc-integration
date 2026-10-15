"""
Gera relatório Excel do sync RM → GEDOC.
"""
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from db_log import init_db, salvar_log_gedoc

# ── paleta ────────────────────────────────────────────────────────────────────
C_RED_DARK  = "C0392B"
C_RED_MID   = "E74C3C"
C_RED_LIGHT = "FADBD8"
C_OK        = "27AE60"
C_HEADER    = "2C3E50"
C_WHITE     = "FFFFFF"
C_GRAY      = "F2F3F4"

_ACAO_FILL = {
    "Criado":     "D5F5E3",
    "Atualizado": "D6EAF8",
    "Desligado":  "FDEBD0",
    "Afastado":   "FCF3CF",
    "Ignorado":   "F2F3F4",
}


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color="000000", size=11) -> Font:
    return Font(bold=bold, color=color, size=size)


def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


class SyncReport:
    _ERROS = {"Erro"}

    _COLS = [
        ("CPF",      "cpf",      16),
        ("Nome",     "nome",     38),
        ("Tipo",     "tipo",     12),
        ("Situação", "situacao", 12),
        ("Ação",     "acao",     14),
        ("Detalhe",  "detalhe",  55),
        ("Hora",     "hora",     10),
    ]

    def __init__(self):
        self._records: list[dict] = []
        self._started_at = datetime.now()
        try:
            init_db()
        except Exception:
            pass

    def record(
        self,
        cpf: str,
        nome: str,
        tipo: str,      # "Interno" | "Externo"
        situacao: str,  # "Ativo" | "Demitido"
        acao: str,      # "Criado" | "Atualizado" | "Desligado" | "Ignorado" | "Erro"
        detalhe: str = "",
    ):
        now = datetime.now()
        self._records.append({
            "cpf":      cpf,
            "nome":     nome,
            "tipo":     tipo,
            "situacao": situacao,
            "acao":     acao,
            "detalhe":  (detalhe or "")[:300],
            "hora":     now.strftime("%H:%M:%S"),
        })
        if not (acao == "Ignorado" and "Demitido" in (detalhe or "")):
            try:
                salvar_log_gedoc(now, cpf, nome, tipo, situacao, acao, (detalhe or "")[:300])
            except Exception:
                pass

    def save(self, output_dir: str = ".") -> str:
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        ts   = self._started_at.strftime("%Y%m%d_%H%M%S")
        path = Path(output_dir) / f"sync_report_{ts}.xlsx"

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        erros       = [r for r in self._records if r["acao"] in self._ERROS]
        processados = [r for r in self._records if r["acao"] not in self._ERROS]

        self._build_resumo(wb, erros)
        if erros:
            self._build_erros(wb, erros)
        self._build_processados(wb, processados)

        wb.save(path)
        return str(path)

    # ── aba Resumo ─────────────────────────────────────────────────────────────

    def _build_resumo(self, wb, erros):
        ws = wb.create_sheet("Resumo")
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 16

        # título
        ws.merge_cells("A1:B1")
        ws["A1"].value     = "RELATÓRIO DE SYNC  RM → GEDOC"
        ws["A1"].fill      = _fill(C_HEADER)
        ws["A1"].font      = _font(bold=True, color=C_WHITE, size=14)
        ws["A1"].alignment = _center()
        ws.row_dimensions[1].height = 30

        # data/hora
        ws.merge_cells("A2:B2")
        ws["A2"].value     = f"Gerado em: {self._started_at.strftime('%d/%m/%Y  %H:%M:%S')}"
        ws["A2"].font      = _font(color="666666", size=10)
        ws["A2"].alignment = _center()
        ws.row_dimensions[2].height = 18

        ws.row_dimensions[3].height = 10

        # alerta
        ws.merge_cells("A4:B4")
        alert = ws["A4"]
        ws.row_dimensions[4].height = 40
        if erros:
            alert.value = f"⚠   ATENÇÃO — {len(erros)} ERRO(S) DURANTE O SYNC"
            alert.fill  = _fill(C_RED_MID)
            alert.font  = _font(bold=True, color=C_WHITE, size=13)
        else:
            alert.value = "✓   SYNC CONCLUÍDO SEM ERROS"
            alert.fill  = _fill(C_OK)
            alert.font  = _font(bold=True, color=C_WHITE, size=13)
        alert.alignment = _center()

        ws.row_dimensions[5].height = 10

        # stats
        def count(acao):
            return sum(1 for r in self._records if r["acao"] == acao)

        n_erros = len(erros)
        rows = [
            ("Total processado", len(self._records)),
            ("Criados",          count("Criado")),
            ("Atualizados",      count("Atualizado")),
            ("Desligados",        count("Desligado")),
            ("Afastados",        count("Afastado")),
            ("Ignorados",        count("Ignorado")),
            ("Erros",            n_erros),
        ]
        for i, (label, val) in enumerate(rows, start=6):
            ws.row_dimensions[i].height = 22
            ca = ws[f"A{i}"]
            cb = ws[f"B{i}"]
            ca.value     = label
            ca.font      = _font(bold=True)
            ca.alignment = _left()
            cb.value     = val
            cb.alignment = _center()
            if label == "Erros" and val > 0:
                ca.fill = _fill(C_RED_LIGHT)
                cb.fill = _fill(C_RED_LIGHT)
                cb.font = _font(bold=True, color=C_RED_DARK)
            elif i % 2 == 0:
                ca.fill = _fill(C_GRAY)
                cb.fill = _fill(C_GRAY)

    # ── aba Erros ──────────────────────────────────────────────────────────────

    def _build_erros(self, wb, erros):
        ws = wb.create_sheet("⚠ Erros")
        self._write_table(ws, erros, header_fill=C_RED_DARK, row_fill=C_RED_LIGHT)

    # ── aba Processados ────────────────────────────────────────────────────────

    def _build_processados(self, wb, rows):
        ws = wb.create_sheet("Processados")
        self._write_table(ws, rows, header_fill=C_HEADER, row_fill=None)

    # ── tabela genérica ────────────────────────────────────────────────────────

    def _write_table(self, ws, rows, *, header_fill, row_fill):
        n_cols = len(self._COLS)

        for ci, (label, _, width) in enumerate(self._COLS, 1):
            ws.column_dimensions[get_column_letter(ci)].width = width
            cell = ws.cell(row=1, column=ci, value=label)
            cell.fill      = _fill(header_fill)
            cell.font      = _font(bold=True, color=C_WHITE)
            cell.alignment = _center()
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}1"

        for ri, rec in enumerate(rows, start=2):
            # cor da linha: por acao (Processados) ou fixo (Erros)
            fill_hex = row_fill or _ACAO_FILL.get(rec.get("acao", ""), C_GRAY)
            for ci, (_, key, _) in enumerate(self._COLS, 1):
                cell = ws.cell(row=ri, column=ci, value=rec.get(key, ""))
                cell.alignment = _left()
                cell.fill      = _fill(fill_hex)
