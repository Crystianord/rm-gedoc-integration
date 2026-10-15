"""
Busca no RM (via rm_client.py) o CODSITUACAO de cada nome presente em uma planilha
e grava o resultado numa nova coluna / novo arquivo .xlsx.

Uso:
    python check_situacao_planilha.py "C:\\caminho\\diferenca_listas.xlsx"
"""
import sys
import unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill

from rm_client import RMClient


def normalize(name: str) -> str:
    if not name:
        return ""
    n = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    return " ".join(n.strip().upper().split())


def main():
    if len(sys.argv) < 2:
        print("Uso: python check_situacao_planilha.py <arquivo.xlsx>")
        sys.exit(1)

    caminho = sys.argv[1]

    print("Buscando carga completa de funcionários no RM...")
    rm = RMClient()
    funcionarios = rm.get_all_employees()
    print(f"{len(funcionarios)} funcionários retornados pelo RM.")

    # Indexa por nome normalizado. Um nome pode ter mais de um registro (chapas
    # diferentes) — guarda todos para reportar se houver ambiguidade.
    por_nome = {}
    for f in funcionarios:
        nome_norm = normalize(f.get("NOME"))
        por_nome.setdefault(nome_norm, []).append(f)

    wb = openpyxl.load_workbook(caminho)
    ws = wb["Ausentes na Lista 2"] if "Ausentes na Lista 2" in wb.sheetnames else wb.active

    header_col = ws.max_column + 1
    ws.cell(row=1, column=header_col, value="CODSITUACAO (RM)").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=header_col).fill = PatternFill("solid", fgColor="C00000")

    col_obs = header_col + 1
    ws.cell(row=1, column=col_obs, value="Observação").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=col_obs).fill = PatternFill("solid", fgColor="C00000")

    ws.column_dimensions[ws.cell(row=1, column=header_col).column_letter].width = 18
    ws.column_dimensions[ws.cell(row=1, column=col_obs).column_letter].width = 30

    nao_encontrados = 0
    ambiguos = 0

    for row in range(2, ws.max_row + 1):
        nome_cell = ws.cell(row=row, column=1)
        nome = nome_cell.value
        if not nome:
            continue

        nome_norm = normalize(nome)
        registros = por_nome.get(nome_norm)

        if not registros:
            ws.cell(row=row, column=header_col, value="NÃO ENCONTRADO")
            nao_encontrados += 1
            continue

        situacoes = {r.get("CODSITUACAO") for r in registros}

        if len(registros) > 1 and len(situacoes) > 1:
            ws.cell(row=row, column=header_col, value=" / ".join(sorted(s or "?" for s in situacoes)))
            ws.cell(row=row, column=col_obs, value=f"{len(registros)} registros com situações diferentes (possível homônimo)")
            ambiguos += 1
        else:
            situacao = registros[0].get("CODSITUACAO")
            ws.cell(row=row, column=header_col, value=situacao or "(vazio)")
            if len(registros) > 1:
                ws.cell(row=row, column=col_obs, value=f"{len(registros)} registros no RM (mesma situação)")

    saida = caminho
    try:
        wb.save(saida)
    except PermissionError:
        saida = caminho.replace(".xlsx", "_com_situacao.xlsx")
        wb.save(saida)
        print(f"AVISO: arquivo original estava aberto/travado. Salvo em novo arquivo: {saida}")

    print(f"Concluído. Não encontrados: {nao_encontrados}. Ambíguos: {ambiguos}.")
    print(f"Arquivo atualizado: {saida}")


if __name__ == "__main__":
    main()
